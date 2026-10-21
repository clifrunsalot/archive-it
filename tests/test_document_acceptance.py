"""Synthetic retrieval/answer acceptance cases; no personal catalogue access."""

import io
import json
import os
import re
import textwrap
from dataclasses import dataclass

import pytest
from docx import Document
from langchain_core.embeddings import Embeddings
from langchain_core.messages import AIMessage
from langchain_core.runnables import RunnableLambda
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from chatbot_tech_doc import get_embeddings, get_llm
from document_library import DocumentLibrary
from document_search import AnswerResult, DocumentSearch, prompt_evidence


@dataclass(frozen=True)
class Case:
    name: str
    question: str
    paragraphs: tuple[str, ...]
    evidence: tuple[str, ...]
    claims: tuple[tuple[str, str], ...]
    forbidden: tuple[str, ...] = ()
    missing: bool = False
    qualifiers: tuple[tuple[str, str], ...] = ()


CONTACT = (
    "Landlord or Agent's Name: Example Rentals. No mailing address is recorded here.",
    "MEGAN'S LAW STATEMENT: Residents should review the public disclosure notice.",
    "2. LANDLORD:\nExample Rentals\n123 Example Street\nExample City, NJ 00000",
    "PROPERTY:\nThe tenant's residence is 987 Tenant Avenue, Other City, NJ 11111.",
)
ADDRESS_CASES = tuple(Case(
    name, question, CONTACT, ("123 Example Street",),
    ((r"123 example street", "123 Example Street"),), ("987 Tenant Avenue",),
) for name, question in (
    ("address-possessive", "What is the landlord's address?"),
    ("address-curly", "What is the landlord\u2019s address?"),
    ("address-paraphrase", "What is the address of the landlord?"),
))
CASES = (*ADDRESS_CASES,
    Case(
        "specific-utilities", "Which specific utilities is the tenant responsible for paying?",
        (
            "UTILITY OBLIGATION: The tenant pays utility charges while responsible for rent.",
            "15. UTILITIES AND SERVICES:\n"
            "Before moving in, transfer gas and electric accounts into your name. "
            "Keep proof of the transfer and submit the proof to the office before collecting keys. "
            "The tenant must pay electric, gas, water, sewer, Internet, telephone and cable television.",
            "OPTIONAL VEHICLE CHARGING: Only residents using the optional charger pay its charging fee.",
        ),
        ("electric, gas, water, sewer, Internet, telephone and cable television",),
        tuple((term, term) for term in ("electric", "gas", "water", "sewer", "internet", "telephone", "cable")),
        qualifiers=((r"charging fee", r"(?:only|using|who use|if|when).*(?:optional|charg)"),),
    ),
    Case(
        "term-versus-signature", "When does the lease term begin and end?",
        (
            "SIGNATURE DETAILS: Signed 2026-04-17. The signing date is not the term start.",
            "TERM: The lease begins 2026-08-01 and ends 2027-07-31.",
            "NOTICE DATES: Notices may be delivered before the end of the term.",
        ),
        ("2026-08-01", "2027-07-31"),
        ((r"2026-08-01", "2026-08-01"), (r"2027-07-31", "2027-07-31")),
        ("2026-04-17",),
    ),
    Case(
        "conditional-payment", "Under what condition does the resident pay the vehicle charging fee, "
        "and what amount is charged per use?",
        (
            "VEHICLE CHARGING FEE:\nOnly residents who use the optional vehicle charger "
            "must pay a charging fee of USD 12 per use. Residents who do not use it owe no charging fee.",
            "GENERAL PAYMENT: Rent is paid monthly. This is separate from optional activities.",
        ),
        ("Only residents who use the optional vehicle charger", "USD 12"),
        ((r"(?:only|if|when|who use|who uses|must use).*(?:charg|use)", "optional vehicle charger"),
         (r"(?:usd\s*12|\$12|12\s*(?:usd|dollars))", "USD 12"),
         (r"per use|each use|each time", "per use")),
    ),
    Case(
        "conflicting-addendum", "What monthly storage fees are stated in the lease and addendum? "
        "Describe any conflict without assuming which one controls.",
        (
            "LEASE STORAGE FEE:\nThe monthly storage fee is USD 40.",
            "ADDENDUM STORAGE FEE:\nThe monthly storage fee is USD 65. "
            "This addendum does not state whether it replaces the lease storage fee.",
        ),
        ("USD 40", "USD 65", "does not state whether it replaces"),
        ((r"(?:usd\s*40|\$40|40\s*(?:usd|dollars))", "USD 40"),
         (r"(?:usd\s*65|\$65|65\s*(?:usd|dollars))", "USD 65"),
         (r"conflict|differ|discrepan|unclear|does not (?:explicitly )?(?:state|specify|clarify)|not (?:state|specif)",
          "does not state whether it replaces")),
        ("supersedes the lease", "addendum controls"),
    ),
    Case(
        "missing-evidence", "What is the swimming pool opening time?",
        (
            "AMENITIES: A fitness room is available to residents.",
            "OFFICE HOURS: The management office opens at 09:00. Office hours are not amenity hours.",
        ),
        (), (), ("09:00", "9:00"), True,
    ),
    Case(
        "nonlease-contact", "What is the supplier's address?",
        (
            "Supplier or Agent's Name: Example Supplies. No mailing address is recorded here.",
            "SUPPLIER NOTICE: Notify the supplier when your delivery address changes.",
            "SUPPLIER:\nExample Supplies\n246 Commerce Road\nExample City, NJ 22222",
            "DELIVERY LOCATION:\n987 Tenant Avenue, Other City, NJ 11111",
        ),
        ("246 Commerce Road",), ((r"246 commerce road", "246 Commerce Road"),),
        ("987 Tenant Avenue",),
    ),
)


class FlatEmbeddings(Embeddings):
    """Intentionally uninformative vectors exercise keyword/heading ranking."""

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[1.0, 0.0] for _ in texts]

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]


def write_document(path, paragraphs):
    if path.suffix == ".txt":
        path.write_text("\n\n".join(paragraphs), encoding="utf-8")
    elif path.suffix == ".docx":
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(str(path))
    else:
        writer = PdfWriter()
        for paragraph in paragraphs:
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
            })
            lines = [line for original in paragraph.splitlines()
                     for line in textwrap.wrap(original, width=75)]
            commands = ["BT /F1 10 Tf 14 TL 50 740 Td"]
            for line in lines:
                escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                commands.append(f"({escaped}) Tj T*")
            commands.append("ET")
            stream = DecodedStreamObject()
            stream.set_data("\n".join(commands).encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        path.write_bytes(output.getvalue())


@pytest.fixture(scope="module", params=["txt", "docx", "pdf"])
def corpus(request, tmp_path_factory):
    folder = tmp_path_factory.mktemp("acceptance-" + request.param)
    library = DocumentLibrary(folder / "catalogue.sqlite3")
    identifiers = {}
    for case in CASES:
        path = folder / f"{case.name}.{request.param}"
        write_document(path, case.paragraphs)
        identifiers[case.name] = library.ingest(path)
    # Deliberately relevant-looking evidence outside every selected scope.
    path = folder / f"unselected.{request.param}"
    write_document(path, ("LANDLORD:\nWrong Organization\n999 Wrong Road",
                         "UTILITIES: The tenant pays no utilities. Pool opens at 06:00."))
    library.ingest(path)
    return library, identifiers


@pytest.fixture(scope="module")
def deterministic_engine(corpus):
    library, _ = corpus
    engine = DocumentSearch(library, FlatEmbeddings(), "acceptance-flat-v1")
    engine.index()
    return engine


@pytest.fixture(scope="module")
def local_engine(corpus):
    library, _ = corpus
    engine = DocumentSearch(library, get_embeddings("nomic-embed-text"), "nomic-embed-text")
    engine.index(rebuild=True)
    return engine


def normalized(text):
    return " ".join(text.casefold().split())


def check_evidence(case, records, library, identifier):
    assert records, f"{case.name}: no evidence supplied"
    source = library.snapshots([identifier])[0]
    for record in records:
        assert record["document_id"] == identifier, f"{case.name}: scope leaked"
        assert record["version_id"] == source.version_id
        segment = source.extraction["segments"][record["segment"] - 1]
        location = record["location"]
        assert segment["text"][location["char_start"]:location["char_end"]] == record["text"]
        assert record["source_path"] == source.source_path
        assert record["quality"]["issues"] == source.extraction["issues"]
    text = normalized(" ".join(record["text"] for record in records))
    for expected in case.evidence:
        assert normalized(expected) in text, f"{case.name}: missing packed evidence: {expected}"


@pytest.mark.parametrize("case", CASES, ids=lambda case: case.name)
def test_retrieval_and_prompt_packing(corpus, deterministic_engine, case):
    library, identifiers = corpus
    identifier = identifiers[case.name]
    captured = []

    def inspect(prompt):
        message = prompt.to_messages()[-1].text
        captured.extend(json.loads(message.split("Evidence (JSON records):\n", 1)[1]
                                   .split("\n\nQuestion:\n", 1)[0]))
        return AIMessage(content='{"status":"no_relevant_evidence","findings":[]}')

    # Checks the actual bounded prompt, not just the larger search result set.
    result = deterministic_engine.answer(case.question, RunnableLambda(inspect), [identifier])
    check_evidence(case, result.evidence, library, identifier)
    assert captured == json.loads(prompt_evidence(result.evidence))
    packed_text = normalized(" ".join(record["text"] for record in captured))
    for expected in case.evidence:
        assert normalized(expected) in packed_text, f"{case.name}: missing actual prompt evidence"
    assert result.coverage["evidence_passages"] == len(captured)
    assert not result.coverage["exhaustive"]
    assert result.coverage["model_requests"] == 1


def check_answer(case, result):
    if case.missing:
        assert result.coverage["insufficient_evidence"], f"{case.name}: invented an answer: {result.answer}"
        assert "not proof of absence" in result.answer
    else:
        assert not result.coverage["insufficient_evidence"], f"{case.name}: failed to answer"
    check_claims(case, result)


def check_claims(case, result):
    if not case.missing:
        records = {record["label"]: record for record in result.evidence}
        for pattern, support in case.claims:
            matching = [line for line in result.answer.splitlines()
                        if re.search(pattern, line, re.IGNORECASE)]
            assert matching, f"{case.name}: missing claim/qualification {pattern}: {result.answer}"
            assert any(
                normalized(support) in normalized(" ".join(records[int(label)]["text"] for label in
                    re.findall(r"\[(\d+)\]", line)))
                for line in matching
            ), f"{case.name}: claim cites no matching evidence: {pattern}: {result.answer!r}"
    for forbidden in case.forbidden:
        assert normalized(forbidden) not in normalized(result.answer), (
            f"{case.name}: distractor/unsupported conclusion: {result.answer}")
    for trigger, qualification in case.qualifiers:
        for line in result.answer.splitlines():
            if re.search(trigger, line, re.IGNORECASE):
                assert re.search(qualification, line, re.IGNORECASE), (
                    f"{case.name}: conditional claim lost its scope: {line}")


@pytest.mark.parametrize(("case_name", "answer"), [
    ("address-possessive", "123 Example Street [2]"),
    ("conditional-payment", "Every resident pays USD 12 per use. [1]"),
    ("conditional-payment", "Only residents who use the optional charger must pay. [1]"),
    ("conflicting-addendum", "The lease charges USD 40; the addendum controls at USD 65. [1]"),
    ("conflicting-addendum", "The lease charges USD 40; the addendum charges USD 65. [1]"),
    ("specific-utilities", "Pay electric, gas, water, sewer, Internet, telephone and cable. [1]\n"
     "All tenants pay the charging fee. [1]"),
])
def test_acceptance_oracle_rejects_missing_qualifications_and_wrong_support(case_name, answer):
    case = next(case for case in CASES if case.name == case_name)
    records = [
        {"label": 1, "text": " ".join(case.paragraphs)},
        {"label": 2, "text": "Unrelated disclosure."},
    ]
    with pytest.raises(AssertionError):
        check_answer(case, AnswerResult(answer, records, {"insufficient_evidence": False}))


def test_missing_evidence_oracle_rejects_invented_hours():
    case = next(case for case in CASES if case.missing)
    with pytest.raises(AssertionError):
        check_answer(case, AnswerResult("The pool opens at 09:00. [1]", [],
                                       {"insufficient_evidence": False}))


@pytest.mark.parametrize("wording", [
    "does not clarify if it replaces",
    "does not clarify whether it replaces",
])
def test_conflict_oracle_accepts_equivalent_explicit_precedence_uncertainty(wording):
    case = next(case for case in CASES if case.name == "conflicting-addendum")
    records = [{"label": 1, "text": case.paragraphs[0]},
               {"label": 2, "text": case.paragraphs[1]}]
    answer = ("The lease states USD 40 monthly. [1]\n"
              f"The addendum states USD 65 monthly and {wording} the lease fee. [2]")
    check_answer(case, AnswerResult(answer, records, {"insufficient_evidence": False}))


def test_multiline_contact_finding_retains_claim_linked_citation_after_rendering():
    from document_answers import parse_focused_findings

    case = ADDRESS_CASES[1]
    records = [{"label": 1, "text": case.paragraphs[2]}]
    response = json.dumps({"status": "findings", "findings": [{
        "support": [1],
        "text": "The landlord's address is:\n123 Example Street\nExample City, NJ 00000",
    }]})
    answer = parse_focused_findings(response, records, 12000)
    check_answer(case, AnswerResult(answer, records, {"insufficient_evidence": False}))
    assert answer.endswith("[1]") and len(answer.splitlines()) == 1


def live_repetitions():
    raw = os.environ.get("RAG_ACCEPTANCE_REPEATS", "1")
    try:
        repeats = int(raw)
    except ValueError as error:
        raise pytest.UsageError("RAG_ACCEPTANCE_REPEATS must be a positive integer") from error
    if repeats < 1:
        raise pytest.UsageError("RAG_ACCEPTANCE_REPEATS must be a positive integer")
    return range(1, repeats + 1)


@pytest.mark.skipif(os.environ.get("RUN_OLLAMA_TESTS") != "1", reason="Enable local RAG acceptance benchmark")
@pytest.mark.parametrize("repeat", live_repetitions(), ids=lambda repeat: f"run-{repeat}")
@pytest.mark.parametrize("case", CASES, ids=lambda case: case.name)
def test_live_retrieval_and_answer(corpus, local_engine, case, repeat, record_property):
    library, identifiers = corpus
    identifier = identifiers[case.name]
    model_name = os.environ.get("RAG_ACCEPTANCE_MODEL", "qwen3.5:4b")
    record_property("answer_model", model_name)
    result = local_engine.answer(case.question, get_llm(model_name), [identifier])
    record_property("case", case.name)
    record_property("repeat", repeat)
    record_property("evidence_passages", len(result.evidence))
    record_property("model_requests", result.coverage["model_requests"])
    check_evidence(case, result.evidence, library, identifier)
    check_answer(case, result)
    assert not result.coverage["exhaustive"]
    assert result.coverage["model_requests"] in (1, 2)
