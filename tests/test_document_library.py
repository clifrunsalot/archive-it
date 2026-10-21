import io
import json
import sqlite3
import sys
import zipfile
from unittest.mock import Mock

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, NumberObject

import document_library as library
from app_logging import LOGGER


@pytest.fixture(autouse=True)
def restore_logging():
    previous, level = LOGGER.handlers[:], LOGGER.level
    LOGGER.handlers = []
    yield
    for handler in LOGGER.handlers:
        handler.close()
    LOGGER.handlers = previous
    LOGGER.setLevel(level)


def make_pdf(*texts: str) -> bytes:
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
            })
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Lease terms", 1)
    document.add_paragraph("Rent: $1,234.50. Due 2026-10-01.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Rent"
    table.cell(1, 1).text = "$1,234.50"
    document.add_paragraph("After table")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_positioned_pdf(page_commands: list[str]) -> bytes:
    writer = PdfWriter()
    for commands in page_commands:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        content = DecodedStreamObject()
        content.set_data(commands.encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_text_preserves_values_paragraphs_and_line_locations():
    result = library.extract_document(
        b"\xef\xbb\xbfPolicy P-123\nDeductible: $500.00\n\nExpires 2027-01-01\n", ".txt",
    )
    assert result.status == "extracted"
    assert not result.ocr_used
    assert result.segments[0].text == "Policy P-123\nDeductible: $500.00"
    assert result.segments[0].location == {"line_start": 1, "line_end": 2}
    assert result.segments[1].location == {"line_start": 4, "line_end": 4}


def test_pdf_text_page_provenance_and_blank_page_quality():
    result = library.safely_extract(make_pdf("Tax year 2025 Total $120.50", "", "Page three"), ".pdf")
    assert result.status == "partial"
    assert result.segments[0].location == {"page": 1}
    assert result.segments[1].location == {"page": 3}
    assert "$120.50" in result.segments[0].text
    assert "page_2_no_extractable_text" in result.issues
    assert "pdf_layout_and_tables_not_verified" in result.issues


def test_docx_preserves_body_order_table_cells_and_heading():
    result = library.safely_extract(make_docx(), ".docx")
    assert [segment.kind for segment in result.segments] == [
        "paragraph", "paragraph", "table_row", "table_row", "paragraph",
    ]
    assert result.segments[3].text == '["Rent", "$1,234.50"]'
    assert result.segments[3].location == {"block": 3, "row": 2, "heading": "Lease terms"}
    assert result.segments[-1].text == "After table"
    assert result.issues == ["docx_textboxes_revisions_nested_tables_not_verified"]


def test_docx_extracts_header_footer_and_first_page_content_with_provenance():
    from docx.enum.section import WD_SECTION
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("Body: monthly rent is USD 1200.")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header: agreement 2027-A."
    table = section.header.add_table(rows=1, cols=2, width=Inches(6))
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    section.footer.paragraphs[0].text = "Footer: page terms."
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "First-page header: effective 2027-01-01."
    section.first_page_footer.paragraphs[0].text = "First-page footer: signed by parties."
    document.settings.odd_and_even_pages_header_footer = True
    section.even_page_header.paragraphs[0].text = "Even-page header: policy 2027-B."
    section.even_page_footer.paragraphs[0].text = "Even-page footer: terms continue."
    second = document.add_section(WD_SECTION.NEW_PAGE)
    # The second section links to the first header/footer; linked text must not duplicate.
    assert second.header.is_linked_to_previous and second.footer.is_linked_to_previous
    output = io.BytesIO()
    document.save(output)

    result = library.safely_extract(output.getvalue(), ".docx")
    by_text = {segment.text: segment for segment in result.segments}
    assert "Body: monthly rent is USD 1200." in by_text
    assert by_text["Header: agreement 2027-A."].location == {
        "section": 1, "area": "header", "block": 1,
    }
    assert by_text["Footer: page terms."].location == {
        "section": 1, "area": "footer", "block": 1,
    }
    assert by_text["First-page header: effective 2027-01-01."].location["area"] == "first_page_header"
    assert by_text["First-page footer: signed by parties."].location["area"] == "first_page_footer"
    assert by_text["Even-page header: policy 2027-B."].location["area"] == "even_page_header"
    assert by_text["Even-page footer: terms continue."].location["area"] == "even_page_footer"
    table_row = next(segment for segment in result.segments if segment.kind == "table_row")
    assert json.loads(table_row.text) == ["Field", "Value"]
    assert table_row.location == {
        "section": 1, "area": "header", "block": 2, "row": 1,
    }
    assert sum(segment.text == "Header: agreement 2027-A." for segment in result.segments) == 1
    assert "docx_textboxes_revisions_nested_tables_not_verified" in result.issues


def test_pdf_repeated_running_lines_are_retained_and_require_review():
    source = make_positioned_pdf([
        "BT /F1 12 Tf 72 750 Td (SYNTHETIC LEASE HEADER) Tj T* "
        "(Rent USD 1200 per month) Tj ET",
        "BT /F1 12 Tf 72 750 Td (SYNTHETIC LEASE HEADER) Tj T* "
        "(Deposit USD 500 payable once) Tj ET",
    ])
    result = library.safely_extract(source, ".pdf")
    assert len(result.segments) == 2
    assert "SYNTHETIC LEASE HEADER" in result.segments[0].text
    assert "SYNTHETIC LEASE HEADER" in result.segments[1].text
    assert "Rent USD 1200" in result.segments[0].text
    assert "Deposit USD 500" in result.segments[1].text
    assert result.status == "partial"
    assert "pdf_repeated_page_lines_requires_review" in result.issues
    assert [segment.location for segment in result.segments] == [{"page": 1}, {"page": 2}]


def test_pdf_column_values_are_preserved_with_layout_limitation_flag():
    source = make_positioned_pdf([
        "BT /F1 12 Tf 72 720 Td (Rent) Tj ET "
        "BT /F1 12 Tf 300 720 Td (USD 1200 monthly) Tj ET "
        "BT /F1 12 Tf 72 690 Td (Deposit) Tj ET "
        "BT /F1 12 Tf 300 690 Td (USD 500 payable once) Tj ET",
    ])
    result = library.safely_extract(source, ".pdf")
    text = result.segments[0].text
    for value in ("Rent", "USD 1200 monthly", "Deposit", "USD 500 payable once"):
        assert value in text
    assert result.segments[0].location == {"page": 1}
    assert "pdf_layout_and_tables_not_verified" in result.issues


def test_extractor_v3_reprocesses_unchanged_docx_for_header_footer_support(tmp_path, monkeypatch):
    source = tmp_path / "lease.docx"
    source.write_bytes(make_docx())
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    current_profile = library.extractor_version(".docx")
    assert current_profile.startswith("extraction-v3:.docx:")
    monkeypatch.setattr(library, "extractor_version", lambda _: "extraction-v2:legacy")
    identifier = store.ingest(source)
    old = store.show(identifier)
    monkeypatch.setattr(library, "extractor_version", lambda _: current_profile)

    store.ingest(source)
    current = store.show(identifier)
    assert current["sha256"] == old["sha256"]
    assert current["id"] != old["id"]
    assert current["extractor"] == current_profile
    assert len(store.versions(identifier)) == 2




@pytest.mark.parametrize("data,suffix,reason", [
    (b"\xff", ".txt", "invalid_utf8"),
    (b"\x00", ".txt", "binary_text"),
    (b" \n", ".txt", "no_extractable_text"),
    (b"fake pdf", ".pdf", "invalid_pdf_signature"),
    (b"text", ".doc", "unsupported_format"),
])
def test_rejected_inputs(data, suffix, reason):
    with pytest.raises(library.ExtractionError, match=reason):
        library.extract_document(data, suffix)


def test_empty_pdf_requires_ocr_or_review():
    with pytest.raises(library.ExtractionError, match="ocr_or_review_required"):
        library.safely_extract(make_pdf(""), ".pdf")


def test_encrypted_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(library.ExtractionError, match="encrypted_pdf"):
        library.safely_extract(output.getvalue(), ".pdf")


def test_extraction_limits(monkeypatch):
    monkeypatch.setattr(library, "MAX_TEXT_CHARS", 3)
    with pytest.raises(library.ExtractionError, match="extracted_text_limit"):
        library.extract_document(b"1234", ".txt")
    monkeypatch.setattr(library, "MAX_PAGES", 1)
    with pytest.raises(library.ExtractionError, match="page_limit"):
        library.safely_extract(make_pdf("", ""), ".pdf")
    monkeypatch.setattr(library, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(library.ExtractionError, match="docx_expansion_limit"):
        library.safely_extract(make_docx(), ".docx")


def test_versions_restart_unchanged_reversion_and_original_export(tmp_path, monkeypatch):
    database = tmp_path / "catalogue.sqlite3"
    source = tmp_path / "lease.txt"
    source.write_bytes(b"Version one $100")
    store = library.DocumentLibrary(database)
    identifier = store.ingest(source)
    first = store.show(identifier)
    source.write_bytes(b"Version two $200")
    assert store.ingest(source) == identifier
    second = store.show(identifier)
    assert second["id"] != first["id"]
    assert len(store.versions(identifier)) == 2
    reopened = library.DocumentLibrary(database)
    assert reopened.show(identifier) == second
    extract = Mock(side_effect=AssertionError("Must not re-extract"))
    monkeypatch.setattr(library, "safely_extract", extract)
    reopened.ingest(source)
    source.write_bytes(b"Version one $100")
    reopened.ingest(source)
    assert reopened.show(identifier)["id"] == first["id"]
    extract.assert_not_called()
    export = tmp_path / "export.txt"
    reopened.export_original(identifier, export, second["id"])
    assert export.read_bytes() == b"Version two $200"
    with pytest.raises(FileExistsError):
        reopened.export_original(identifier, export)


def test_failed_updates_keep_previous_version_and_record_status(tmp_path):
    source = tmp_path / "record.txt"
    source.write_bytes(b"Receipt $30.00")
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    identifier = store.ingest(source)
    previous = store.show(identifier)
    source.write_bytes(b"\xff")
    with pytest.raises(library.ExtractionError):
        store.ingest(source)
    assert store.show(identifier) == previous
    entry = store.list_documents()[0]
    assert entry["latest_attempt_status"] == "failed"
    assert entry["latest_attempt_reason"] == "invalid_utf8"
    assert entry["quality"]["status"] == "extracted"
    source.unlink()
    assert store.show(identifier) == previous
    with pytest.raises(library.ExtractionError):
        store.ingest(source)
    assert len(store.versions(identifier)) == 1


@pytest.mark.parametrize("filename,data", [
    ("broken.pdf", b"%PDF-invalid"),
    ("broken.docx", b"not zip"),
    ("unsupported.doc", b"old Word"),
])
def test_failed_first_ingestion_is_catalogued(tmp_path, filename, data):
    source = tmp_path / filename
    source.write_bytes(data)
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    with pytest.raises(library.ExtractionError):
        store.ingest(source)
    entry = store.list_documents()[0]
    assert entry["current_version"] is None
    assert entry["quality"] is None
    assert entry["latest_attempt_status"] == "failed"


def test_file_size_limit_and_distinct_same_name_documents(tmp_path, monkeypatch):
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    for folder in ("a", "b"):
        source = tmp_path / folder / "receipt.txt"
        source.parent.mkdir()
        source.write_bytes(b"Total $12")
        store.ingest(source)
    assert len(store.list_documents()) == 2
    monkeypatch.setattr(library, "MAX_FILE_BYTES", 3)
    with pytest.raises(library.ExtractionError, match="file_size_limit"):
        store.ingest(source)


def test_docx_images_and_macros_are_explicit():
    original = make_docx()
    for name, expected in [
        ("word/media/image.png", "images_not_extracted"),
        ("word/vbaProject.bin", "macro_document_not_supported"),
    ]:
        output = io.BytesIO(original)
        with zipfile.ZipFile(output, "a") as archive:
            archive.writestr(name, b"image or macro")
        if "vba" in name:
            with pytest.raises(library.ExtractionError, match=expected):
                library.safely_extract(output.getvalue(), ".docx")
        else:
            result = library.safely_extract(output.getvalue(), ".docx")
            assert result.status == "partial"
            assert expected in result.issues


def test_parser_diagnostics_are_redacted(tmp_path, monkeypatch):
    import logging
    from app_logging import configure_logging, LOGGER

    config = tmp_path / "logging.json"
    config.write_text(json.dumps({
        "level": "DEBUG", "file": "log.txt", "max_bytes": 5000, "retained_files": 5,
    }))
    previous, level = LOGGER.handlers[:], LOGGER.level
    configure_logging(config)
    try:
        def extract(_data, _suffix):
            logging.getLogger("pypdf").warning("SECRET-DOCUMENT-TEXT")
            return library.Extraction([library.Segment("SECRET", {"page": 1})], [], "extracted")
        monkeypatch.setattr(library, "extract_document", extract)
        result = library.safely_extract(b"unused", ".pdf")
        assert result.status == "partial"
        assert "parser_diagnostics_require_review" in result.issues
        assert "SECRET" not in (tmp_path / "log.txt").read_text()
    finally:
        for handler in LOGGER.handlers:
            handler.close()
        LOGGER.handlers = previous
        LOGGER.setLevel(level)


def test_cli_batch_failure_status_and_json_output(tmp_path, monkeypatch, capsys):
    config = tmp_path / "logging.json"
    config.write_text(json.dumps({
        "level": "INFO", "file": "log.txt", "max_bytes": 5000, "retained_files": 5,
    }))
    source = tmp_path / "sample.txt"
    source.write_text("Synthetic receipt $5")
    bad = tmp_path / "old.doc"
    bad.write_text("unsupported")
    prefix = ["document_library.py", "--database", str(tmp_path / "catalogue.sqlite3"),
              "--log-config", str(config)]
    monkeypatch.setattr("sys.argv", prefix + ["ingest", str(source), str(bad)])
    assert library.main() == 1
    output = capsys.readouterr()
    identifier = json.loads(output.out)["document_id"]
    assert "unsupported_format" in output.err
    monkeypatch.setattr("sys.argv", prefix + ["show", identifier])
    assert library.main() == 0
    assert json.loads(capsys.readouterr().out)["extraction"]["segments"][0]["text"] == "Synthetic receipt $5"


def test_enumerate_cli_command_is_registered(monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["document_library.py", "enumerate", "--help"])
    with pytest.raises(SystemExit) as result:
        library.main()
    assert result.value.code == 0
    assert "question" in capsys.readouterr().out


def test_unknown_schema_and_wrong_document_version(tmp_path):
    database = tmp_path / "unknown.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.execute("PRAGMA user_version=999")
    with pytest.raises(ValueError, match="schema"):
        library.DocumentLibrary(database)
    store = library.DocumentLibrary(tmp_path / "valid.sqlite3")
    with pytest.raises(ValueError, match="not found"):
        store.show("nonexistent")


def test_storage_failure_rolls_back_new_version_and_current_pointer(tmp_path):
    source = tmp_path / "receipt.txt"
    source.write_text("Total $10")
    database = tmp_path / "catalogue.sqlite3"
    store = library.DocumentLibrary(database)
    identifier = store.ingest(source)
    previous = store.show(identifier)
    with sqlite3.connect(database) as connection:
        connection.execute("""
            CREATE TRIGGER fail_publication BEFORE UPDATE OF current_version ON documents
            BEGIN SELECT RAISE(ABORT, 'Synthetic publication failure'); END
        """)
    source.write_text("Total $20")
    with pytest.raises(sqlite3.IntegrityError):
        store.ingest(source)
    assert library.DocumentLibrary(database).show(identifier) == previous
    assert len(store.versions(identifier)) == 1


def test_parser_version_change_creates_new_snapshot(tmp_path, monkeypatch):
    source = tmp_path / "receipt.txt"
    source.write_text("Total $10")
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    identifier = store.ingest(source)
    monkeypatch.setattr(library, "extractor_version", lambda _: "extraction-v2")
    store.ingest(source)
    assert len(store.versions(identifier)) == 2
    assert store.show(identifier)["extractor"] == "extraction-v2"


def test_cli_persists_across_processes(tmp_path):
    import subprocess
    import sys

    source = tmp_path / "receipt.txt"
    source.write_text("Synthetic receipt $5")
    config = tmp_path / "logging.json"
    config.write_text(json.dumps({
        "level": "DEBUG", "file": "log.txt", "max_bytes": 5000, "retained_files": 5,
    }))
    command = [
        sys.executable, str(library.PROJECT_DIR / "document_library.py"),
        "--database", str(tmp_path / "catalogue.sqlite3"), "--log-config", str(config),
    ]
    ingested = subprocess.run(command + ["ingest", str(source)],
                              check=True, capture_output=True, text=True)
    identifier = json.loads(ingested.stdout)["document_id"]
    source.unlink()
    shown = subprocess.run(command + ["show", identifier],
                           check=True, capture_output=True, text=True)
    assert json.loads(shown.stdout)["extraction"]["segments"][0]["text"] == "Synthetic receipt $5"
    text = (tmp_path / "log.txt").read_text()
    assert "Document version published status=extracted" in text
    assert "Synthetic receipt" not in text
    assert "receipt.txt" not in text


@pytest.mark.parametrize("suffix,contents,location_key", [
    (".pdf", lambda: make_pdf("Receipt R-777 Total $123.45"), "page"),
    (".docx", make_docx, "block"),
])
def test_pdf_docx_search_retains_extraction_provenance(tmp_path, suffix, contents, location_key):
    from document_search import DocumentSearch
    from langchain_core.embeddings import Embeddings

    source = tmp_path / ("document" + suffix)
    source.write_bytes(contents())
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    identifier = store.ingest(source)
    embeddings = Mock(spec=Embeddings)
    embeddings.embed_documents.side_effect = lambda texts: [[1.0, 0.0] for _ in texts]
    embeddings.embed_query.return_value = [1.0, 0.0]
    engine = DocumentSearch(store, embeddings, "synthetic")
    engine.index()
    result = engine.search("Total Rent", [identifier], limit=10)
    assert result["passages"]
    assert all(location_key in passage["location"] for passage in result["passages"])
    assert all(passage["quality"]["issues"] for passage in result["passages"])
    assert all(passage["source"] == source.name for passage in result["passages"])


def test_embedded_pdf_page_bounds_correct_repeated_text(monkeypatch):
    monkeypatch.setattr(library, "MAX_TEXT_CHARS", 50)
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    resources = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    form = DecodedStreamObject()
    form.update({
        NameObject("/Type"): NameObject("/XObject"),
        NameObject("/Subtype"): NameObject("/Form"),
        NameObject("/BBox"): ArrayObject([NumberObject(n) for n in (0, 0, 612, 1584)]),
        NameObject("/Resources"): resources,
    })
    form.set_data(b"BT /F1 12 Tf 72 720 Td (First receipt $10) Tj ET "
                  b"BT /F1 12 Tf 72 1512 Td (Second receipt $20) Tj ET")
    reference = writer._add_object(form)
    for offset in (0, -792):
        page = writer.add_blank_page(612, 792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Full"): reference}),
        })
        content = DecodedStreamObject()
        content.set_data(f"q 1 0 0 1 0 {offset} cm /Full Do Q".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    result = library.safely_extract(output.getvalue(), ".pdf")
    assert "First receipt $10" in result.segments[0].text
    assert "Second" not in result.segments[0].text
    assert "Second receipt $20" in result.segments[1].text
    assert "First" not in result.segments[1].text
    assert result.status == "extracted"
    assert "pdf_page_bounds_reextracted" in result.issues
    assert not result.ocr_used


def test_genuinely_repeated_pages_remain_present_and_flagged():
    result = library.safely_extract(make_pdf("Same receipt", "Same receipt"), ".pdf")
    assert len(result.segments) == 2
    assert result.status == "partial"
    assert "pdf_repeated_page_text_requires_review" in result.issues


@pytest.mark.parametrize("message,status,code", [
    ("Superfluous whitespace found in object header b'1' b'0'",
     "extracted", "pdf_object_header_whitespace"),
    ("Install fontTools to decode", "partial", "pdf_font_decoding_requires_review"),
    ("SECRET unknown parser error", "partial", "parser_diagnostics_require_review"),
])
def test_classified_diagnostics_log_final_status_only(monkeypatch, caplog, message, status, code):
    import logging

    def extract(_data, _suffix):
        logging.getLogger("pypdf").warning(message)
        return library.Extraction([library.Segment("value", {"page": 1})], [], "extracted")
    monkeypatch.setattr(library, "extract_document", extract)
    with caplog.at_level(logging.DEBUG, logger="archive_it"):
        result = library.safely_extract(b"unused", ".pdf")
    assert result.status == status
    assert code in result.issues
    assert f"Document extracted status={status}" in caplog.text
    assert message not in caplog.text


def test_failure_logs_once_with_safe_correlation(tmp_path, caplog):
    import logging

    source = tmp_path / "SECRET.txt"
    source.write_bytes(b"\xff")
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    with caplog.at_level(logging.DEBUG, logger="archive_it"):
        with pytest.raises(library.ExtractionError, match="invalid_utf8"):
            store.ingest(source)
    errors = [r for r in caplog.records if r.levelno == logging.ERROR]
    assert len(errors) == 1
    assert "reason=invalid_utf8" in errors[0].message
    assert "document_id=" in errors[0].message and "attempt_id=" in errors[0].message
    assert "SECRET" not in caplog.text


def test_pdf_uses_separate_file_limit(tmp_path, monkeypatch):
    source = tmp_path / "receipt.pdf"
    source.write_bytes(make_pdf("Total $10"))
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    monkeypatch.setattr(library, "MAX_FILE_BYTES", 3)
    identifier = store.ingest(source)
    assert store.show(identifier)["format"] == ".pdf"
    monkeypatch.setattr(library, "MAX_PDF_FILE_BYTES", 3)
    with pytest.raises(library.ExtractionError, match="file_size_limit"):
        store.ingest(source)


def test_folder_ingestion_skips_system_metadata(tmp_path, monkeypatch, capsys):
    folder = tmp_path / "inputs"
    folder.mkdir()
    (folder / ".DS_Store").write_bytes(b"metadata")
    (folder / "._receipt.txt").write_bytes(b"metadata")
    (folder / "receipt.txt").write_text("Total $10")
    config = tmp_path / "logging.json"
    config.write_text(json.dumps({
        "level": "DEBUG", "file": "log.txt", "max_bytes": 5000, "retained_files": 5,
    }))
    database = tmp_path / "catalogue.sqlite3"
    monkeypatch.setattr("sys.argv", ["document_library.py", "--database", str(database),
                                    "--log-config", str(config), "ingest", str(folder)])
    assert library.main() == 0
    assert len(library.DocumentLibrary(database).list_documents()) == 1
    assert not capsys.readouterr().err


def test_image_ocr_flags_and_search_provenance(tmp_path, monkeypatch):
    from document_search import DocumentSearch
    from langchain_core.embeddings import Embeddings

    monkeypatch.setattr(library, "extract_local",
                        lambda *_: {"pages": [{"text": "Receipt R-1 $12", "ocr": True}]})
    source = tmp_path / "scan.png"
    source.write_bytes(b"synthetic")
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    identifier = store.ingest(source)
    extraction = store.show(identifier)["extraction"]
    assert extraction["status"] == "partial" and extraction["ocr_used"]
    assert extraction["issues"] == ["ocr_text_requires_review"]
    embedding = Mock(spec=Embeddings)
    embedding.embed_documents.side_effect = lambda texts: [[1.0, 0.0] for _ in texts]
    embedding.embed_query.return_value = [1.0, 0.0]
    engine = DocumentSearch(store, embedding, "synthetic")
    engine.index()
    passage = engine.search("Receipt", [identifier])["passages"][0]
    assert passage["location"]["method"] == "ocr_tesseract_eng"
    assert passage["quality"]["ocr_used"]


def test_failed_ocr_preserves_previous_version(tmp_path, monkeypatch):
    source = tmp_path / "scan.png"
    source.write_bytes(b"first")
    monkeypatch.setattr(library, "extract_local",
                        lambda *_: {"pages": [{"text": "Total $10", "ocr": True}]})
    store = library.DocumentLibrary(tmp_path / "catalogue.sqlite3")
    identifier = store.ingest(source)
    first = store.show(identifier)
    source.write_bytes(b"second")
    monkeypatch.setattr(library, "extract_local",
                        Mock(side_effect=library.LocalExtractionError("ocr_timeout")))
    with pytest.raises(library.ExtractionError, match="ocr_timeout"):
        store.ingest(source)
    assert store.show(identifier) == first
    assert store.list_documents()[0]["latest_attempt_reason"] == "ocr_timeout"
